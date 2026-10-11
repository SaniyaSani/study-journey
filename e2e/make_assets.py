"""Creates the small test files the browser tests upload. Needs: pip install reportlab python-pptx"""
import os, wave
H = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets'); os.makedirs(H, exist_ok=True)
w = wave.open(H + '/lecture3.wav', 'wb'); w.setnchannels(1); w.setsampwidth(1); w.setframerate(4000); w.writeframes(bytes([128]) * (4000 * 1400)); w.close()  # 23 min of silence
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
c = canvas.Canvas(H + '/ESC401_Week3_slides.pdf', pagesize=A4)
c.setFont('Helvetica-Bold', 20); c.drawString(60, 780, 'ESC401 High Performance Computing - Week 3'); c.setFont('Helvetica', 14); c.drawString(60, 740, 'Linking and compilation'); c.drawString(60, 715, 'The compiler turns C source into an executable.'); c.showPage()
c.setFont('Helvetica-Bold', 18); c.drawString(60, 780, 'Your first MPI program'); c.setFont('Courier', 13)
for i, l in enumerate(['#include <mpi.h>', 'int main(int argc, char **argv) {', '  MPI_Init(&argc, &argv);', '  MPI_Finalize();', '}']): c.drawString(60, 740 - i * 18, l)
c.setFont('Helvetica', 13); c.drawString(60, 620, 'Every MPI program must call MPI_Init before any other MPI call.'); c.showPage()
c.setFont('Helvetica', 14); c.drawString(60, 780, 'Exercise: run hello_mpi with four processes.'); c.showPage(); c.save()
from pptx import Presentation
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = 'Week 3 slides: compiling'; s.placeholders[1].text = 'gcc hello.c -o hello\nMPI_Init is first'
s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = 'Linking'; s.placeholders[1].text = '-lm goes after the source files'; p.save(H + '/ESC401_week3.pptx')
open(H + '/broken.pdf', 'wb').write(b'%PDF-1.4 this is not really a pdf'); open(H + '/notes.docx', 'wb').write(b'PK\x03\x04fake')
print('assets ready in', H)
